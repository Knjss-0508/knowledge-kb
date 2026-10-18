from __future__ import annotations

import json

import pytest

from answer_hub.cz_integration import CzIntegrationAdapter, CzIntegrationConfig
from answer_hub.mimo import MimoClient, MimoConfig
from answer_hub.workflow import write_topic_candidate_knowledge_workbook


def _client(response, captured):
    client = MimoClient(MimoConfig(
        api_key="test", base_url="https://example.invalid/v1", model="test",
    ))

    def post(payload):
        captured.append(payload)
        return {"choices": [{"message": {"content": json.dumps(response, ensure_ascii=False)}}]}

    client._post = post
    return client


@pytest.mark.parametrize("standards_enabled", [False, True])
def test_transcription_request_allows_one_rule_without_inventing_a_workflow(standards_enabled):
    captured = []
    candidate = {
        "title": "折叠屏手机折叠后关机如何判定",
        "subtitles": [],
        "content": "1. 折叠屏手机在折叠动作后出现自动关机时，按自动关机判定。",
        "recommended_reply": "折叠屏手机在折叠动作后出现自动关机时，按自动关机判定。",
        "content_type": "定义型", "category_l1": "手机", "category_l2": "开机情况",
        "layer": "L2", "knowledge_form": "具体判定", "standard_refs": [],
        "applicable_scope": "手机", "applicable_brands": [], "applicable_models": [],
        "confidence": 0.9, "reasoning_summary": "实际答复支持判自动关机。",
        "needs_human_review": True, "image_evidence_summary": "文字已能表达本条结论。",
        "requires_images": False, "image_usage_instruction": "",
    }
    result = _client(candidate, captured).label_topic(
        {"产品类型": "手机"}, [], use_standard_references=standards_enabled, max_attempts=1,
    )
    prompt = captured[0]["messages"][1]["content"][0]["text"]
    assert "默认沉淀核验过程" not in prompt
    assert "禁止重复整段正文" not in prompt
    assert "正文只有一句完整知识时，推荐回复可以与正文相同" in prompt
    assert "不得因问题涉及图片或功能异常就自动补写核验流程" in prompt
    assert result.candidate["recommended_reply"] == candidate["recommended_reply"]


def test_value_request_does_not_reject_an_explicit_rule_only_for_being_one_sentence():
    captured = []
    classification = {
        "topic_stage": "质检标准", "knowledge_value": "值得沉淀",
        "stage_reason": "明确现象与具体判定选项。",
        "value_reason": "可回答类似问题。",
        "reusable_knowledge": "折叠屏手机折叠后关机按自动关机判定。",
        "confidence": 0.9, "needs_human_review": True,
    }
    _client(classification, captured).classify_topic_stage({"member_count": 1})
    prompt = captured[0]["messages"][1]["content"][0]["text"]
    assert "单成员主题若只是" not in prompt
    assert "只有一句话、只有一个来源或没有额外步骤，都不是单独否决理由" in prompt
    assert "不得把模糊的“正常、异常、可以”自动改写成通用规则" in prompt


def test_content_review_request_preserves_one_sentence_and_equal_reply_policy():
    captured = []
    review = {
        "decision": "需修改", "knowledge_value": "待确认", "error_type": "",
        "reason": "待人工复核", "standard_consistency": "无可信标准",
        "evidence_sufficiency": "充分", "content_consistency": "一致",
        "image_necessity": "不需要", "title_quality": "清晰",
        "confidence": 0.8, "priority_review": True,
    }
    _client(review, captured).review_topic(
        {"knowledge_value": "待确认"}, {"content": "一条完整规则。"}, [],
        use_standard_references=False,
    )
    prompt = captured[0]["messages"][1]["content"][0]["text"]
    assert "正文只有一句或推荐回复与正文相同，不构成质量问题" in prompt
    assert "只有来源明确提供核验步骤时，才能接受流程型知识" in prompt


@pytest.mark.parametrize("standards_enabled", [False, True])
def test_empty_body_is_not_exported_even_with_all_approval_fields(tmp_path, standards_enabled):
    from openpyxl import load_workbook

    candidate = {
        "主题ID": "EMPTY", "主标题": "折叠屏手机折叠后关机如何判定",
        "知识内容": "", "推荐回复": "按自动关机判定。",
        "审核结论": "通过", "主题沉淀价值": "值得沉淀",
        "主题转写状态": "topic_model_labeled", "模型调用状态": "model_success",
        "模型输出校验状态": "passed", "模型质量状态": "passed",
        "知识草稿状态": "ready_for_human_review", "模型初标结论": "通过",
        "关联标准项": "STD-TEST",
    }
    path = tmp_path / "candidate.xlsx"
    write_topic_candidate_knowledge_workbook(
        [candidate], path, use_standard_references=standards_enabled,
    )
    book = load_workbook(path, read_only=True)
    try:
        assert list(book.worksheets[0].iter_rows(min_row=2, values_only=True)) == []
    finally:
        book.close()
    adapter = CzIntegrationAdapter(CzIntegrationConfig("https://example.invalid", "test"))
    assert "缺少知识内容。" in adapter.validate_candidate(
        candidate, "test-category", "empty-case", require_eligible=False,
    )


@pytest.mark.parametrize(
    "interface_location,expected_length",
    [pytest.param("侧面", 100, id="100-characters"),
     pytest.param("左侧面", 101, id="101-characters")],
)
def test_public_build_preserves_an_identical_single_sentence_reply_over_100_characters(
    interface_location, expected_length,
):
    from types import SimpleNamespace

    from answer_hub.mimo import MimoLabelResult
    from answer_hub.workflow import build_topic_review_rows

    sentence = (
        "对于2019款16寸MacBook Pro笔记本电脑，在设备型号已经核实为该款机型"
        f"且当前问题仅询问机身{interface_location}耳机接口的物理规格、"
        "没有涉及耳机无法插入或音频输出故障的情况下，耳机接口规格按3.5mm填写。"
    )
    assert len(sentence) == expected_length

    def label_topic(topic, matches, **kwargs):
        return MimoLabelResult(candidate={
            "title": "2019款16寸MacBook Pro耳机孔是什么规格",
            "subtitles": [], "content": "1. " + sentence,
            "recommended_reply": sentence, "knowledge_form": "具体判定",
            "content_type": "定义型", "confidence": 0.95,
            "standard_refs": [], "applicable_scope": "笔记本",
            "applicable_brands": [],
            "applicable_models": ["2019款16寸MacBook Pro"],
        }, request_audit={}, response_audit={})

    topics, _, gaps, _ = build_topic_review_rows(
        [{
            "数据ID": "SPEC-LONG", "_原子知识ID": "SPEC-LONG-U1",
            "产品类型": "笔记本",
            "核心问题": "2019款16寸MacBook Pro耳机接口规格如何填写",
            "对象/部位": "耳机接口", "异常现象": "接口规格查询",
            "聊天内容": (
                "回收师：2019款16寸MacBook Pro耳机接口规格如何填写\n"
                "答疑：" + sentence
            ),
        }],
        mimo_client=SimpleNamespace(
            config=SimpleNamespace(model="test"), label_topic=label_topic,
        ),
        use_mimo=True, clustering_mode="rule", use_standard_references=False,
        transcribe_all_admitted_topics=True,
    )

    assert not gaps
    assert len(topics) == 1
    topic = topics[0]
    assert topic["知识内容"] == "1. " + sentence
    assert topic["主题无来源内容"] == ""
    assert json.loads(topic["转写质量审计"])["status"] == "passed"
    assert topic["推荐回复"] == sentence
    assert topic["模型质量状态"] == "passed"
    assert topic["知识草稿状态"] == "ready_for_human_review"
