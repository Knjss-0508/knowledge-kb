from pathlib import Path
from concurrent.futures import ThreadPoolExecutor
import json
import threading
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

import answer_hub.topic_registry as topic_registry_module
import answer_hub.workflow as workflow_module
from answer_hub.audit import AuditStore
from answer_hub.cz_integration import CzPublishedKnowledgeItem
from answer_hub.cz_topic_snapshot import CZTopicSnapshot
from answer_hub.mimo import MimoLabelResult
from answer_hub.topic_registry import TopicRegistry
from answer_hub.workflow import build_topic_review_rows, write_cluster_only_workbook


def _row(
    work_order_id: str,
    *,
    product_type: str = "手机",
    business_line: str = "自营回收",
    subject: str = "手机外壳",
    phenomenon: str = "碎裂",
) -> dict[str, str]:
    return {
        "数据ID": work_order_id,
        "工单ID": work_order_id,
        "原始工单ID": work_order_id,
        "回收业务层级": business_line,
        "产品类型": product_type,
        "聊天内容": f"{subject}{phenomenon}应该怎么判定？",
        "核心问题": f"{subject}{phenomenon}如何判定",
        "模型主题一级分类": "外观问题",
        "模型主题二级分类": "外壳外观",
        "问题意图": "标准判定",
        "对象/部位": subject,
        "异常现象": phenomenon,
        "解题方式": "按外观质检口径判定",
        "_原子知识ID": f"{work_order_id}-U1",
    }


def _signature_value(signature: dict[str, object], field: str) -> str:
    values = signature.get(field)
    return values[0] if isinstance(values, list) and values else ""


class _HighConfidenceDirectMimo:
    config = SimpleNamespace(model="incremental-topic-test")

    def analyze_cluster_units(
        self,
        row: dict[str, str],
    ) -> MimoLabelResult:
        return MimoLabelResult(
            candidate={
                "conversation_type": "single_topic",
                "reason": "会话只有一个清晰问题。",
                "topics": [
                    {
                        "normalized_issue": row["核心问题"],
                        "product_category": row["产品类型"],
                        "scope_type": "品类专用",
                        "platform": "通用",
                        "brand": "通用",
                        "model_scope": "通用",
                        "category_l1": row["模型主题一级分类"],
                        "category_l2": row["模型主题二级分类"],
                        "intent": row["问题意图"],
                        "subject": row["对象/部位"],
                        "phenomenon": row["异常现象"],
                        "judgment_target": row["核心问题"],
                        "resolution_mode": row["解题方式"],
                        "standard_path": row["核心问题"],
                        "threshold_or_exception": "无明确阈值",
                        "evidence_summary": row["聊天内容"],
                        "confidence": 0.95,
                        "requires_review": False,
                    }
                ],
            },
            request_audit={},
            response_audit={},
        )

    def cluster_atomic_units(self, _units):
        raise AssertionError("清晰单原子问题不应额外调用聚类模型")


class _ReviewRequiredDirectMimo(_HighConfidenceDirectMimo):
    def analyze_cluster_units(
        self,
        row: dict[str, str],
    ) -> MimoLabelResult:
        result = super().analyze_cluster_units(row)
        result.candidate["topics"][0]["requires_review"] = True
        return result


def _published_business_item(
    knowledge_id: str,
    title: str,
    *,
    score: float = 0.95,
) -> CzPublishedKnowledgeItem:
    return CzPublishedKnowledgeItem(
        knowledge_id=knowledge_id,
        title=title,
        text=title,
        knowledge_origin="business_accumulation",
        business_type="self_operated",
        category_id="cat-notebook",
        level1_label="基本情况",
        product_type="笔记本",
        models=(),
        keywords=(),
        source_ref=f"knowledge-kb://knowledge/{knowledge_id}",
        status="published",
        version="test-v1",
        score=score,
        final_score=score,
    )


def test_cluster_only_reuses_published_business_knowledge_without_local_registration(
    tmp_path: Path,
) -> None:
    audit = AuditStore(tmp_path / "audit.db")
    row = _row(
        "2085229024728059925",
        product_type="笔记本",
        subject="硬盘",
        phenomenon="品牌是否为第三方",
    )
    row.update(
        {
            "核心问题": "笔记本硬盘品牌是否为第三方",
            "判定目标": "确认硬盘是否为品牌件",
            "判定结论": "确认硬盘是否为品牌件",
        }
    )

    def retrieve(*_args):
        return [
            _published_business_item(
                "BIZ-STORAGE-BRAND",
                "笔记本硬盘是否为品牌件",
            )
        ], {"source": "cz_published_knowledge"}

    topics, _mapping, _gaps, _pending = build_topic_review_rows(
        [row],
        use_mimo=False,
        mimo_client=_HighConfidenceDirectMimo(),
        audit_store=audit,
        run_id="run-online-business-auto-reuse",
        clustering_mode="direct_mimo",
        use_standard_references=False,
        cluster_only=True,
        enforce_cluster_admission=True,
        topic_business_knowledge_retriever=retrieve,
    )

    assert topics[0]["线上已有知识匹配结果"] == "auto_reuse"
    assert topics[0]["线上已有知识匹配ID"] == "BIZ-STORAGE-BRAND"
    assert topics[0]["线上已有知识边界状态"] == "boundary_complete"
    assert "storage_brand" in topics[0]["线上已有知识边界标记"]
    assert topics[0]["历史主题处理结果"] == "reused_published_business_knowledge"
    assert audit.list_registered_topics("自营回收", "笔记本") == []


def test_cluster_only_routes_low_confidence_business_knowledge_to_manual_review(
    tmp_path: Path,
) -> None:
    audit = AuditStore(tmp_path / "audit.db")
    row = _row(
        "2085229024728059926",
        product_type="笔记本",
        subject="硬盘",
        phenomenon="品牌是否为第三方",
    )
    row.update(
        {
            "核心问题": "笔记本硬盘品牌是否为第三方",
            "判定目标": "确认硬盘是否为品牌件",
            "判定结论": "确认硬盘是否为品牌件",
        }
    )

    def retrieve(*_args):
        return [
            _published_business_item(
                "BIZ-STORAGE-LOW",
                "笔记本硬盘是否为品牌件",
                score=0.80,
            )
        ], {"source": "cz_published_knowledge"}

    topics, _mapping, _gaps, _pending = build_topic_review_rows(
        [row],
        use_mimo=False,
        mimo_client=_HighConfidenceDirectMimo(),
        audit_store=audit,
        run_id="run-online-business-manual-review",
        clustering_mode="direct_mimo",
        use_standard_references=False,
        cluster_only=True,
        enforce_cluster_admission=True,
        topic_business_knowledge_retriever=retrieve,
    )

    assert topics[0]["线上已有知识匹配结果"] == "manual_review"
    assert topics[0]["线上已有知识边界状态"] == "boundary_complete"
    assert topics[0]["聚类准入状态"] == "待人工线上知识复核"
    assert topics[0]["是否重点复核"] == "是"
    assert audit.list_registered_topics("自营回收", "笔记本") == []


def test_registry_appends_compatible_new_evidence_to_existing_topic(
    tmp_path: Path,
) -> None:
    registry = TopicRegistry(AuditStore(tmp_path / "audit.db"))

    first = registry.integrate(
        proposed_topic_id="TOP-FIRST",
        topic_key=("direct_mimo", "自营回收", "手机", "cluster-1"),
        rows=[_row("WO-001", phenomenon="碎裂")],
        run_id="run-1",
    )
    second = registry.integrate(
        proposed_topic_id="TOP-SECOND",
        topic_key=("direct_mimo", "自营回收", "手机", "cluster-1"),
        rows=[_row("WO-002", phenomenon="掉漆")],
        run_id="run-2",
    )

    assert first.topic_id == "TOP-FIRST"
    assert second.topic_id == "TOP-FIRST"
    assert second.matched_existing is True
    assert second.added_member_count == 1
    assert second.evidence_version == first.evidence_version
    assert second.requires_re_review is True
    assert second.incremental_supplement_pending is True
    assert {
        row["原始工单ID"]
        for row in second.rows
    } == {"WO-001", "WO-002"}


def test_registry_uses_published_local_cz_snapshot_as_trusted_history(
    tmp_path: Path,
) -> None:
    audit = AuditStore(tmp_path / "audit.db")
    historical_row = _row("WO-CZ-SNAPSHOT-BASE", phenomenon="碎裂")
    signature = topic_registry_module._topic_signature([historical_row])
    snapshot = CZTopicSnapshot(tmp_path / "cz-topic-snapshot.db")
    snapshot.sync(
        [
            {
                "cz_topic_id": "CZ-TOPIC-OUTER-HOUSING",
                "status": "published",
                "title": "手机外壳外观损伤判定",
                "business_line": "自营回收",
                "product_category": "手机",
                "source_topic_key": "TOP-CZ-OUTER-HOUSING",
                "standard_family": _signature_value(signature, "standard_families"),
                "merge_policy": _signature_value(signature, "merge_policies"),
                "object_key": _signature_value(signature, "object_keys"),
                "phenomenon_value": _signature_value(signature, "phenomenon_values"),
                "query_target": _signature_value(signature, "query_targets"),
                "detection_target": _signature_value(signature, "detection_targets"),
                "platform": _signature_value(signature, "platforms"),
                "brand": _signature_value(signature, "brands"),
                "model_scope": _signature_value(signature, "model_scopes"),
                "threshold_values": _signature_value(signature, "threshold_values"),
            }
        ],
        snapshot_version="local-cz-v1",
    )

    resolution = TopicRegistry(audit, cz_snapshot=snapshot).integrate(
        proposed_topic_id="TOP-NEW-FROM-CZ",
        topic_key=("direct_mimo", "自营回收", "手机", "cluster-1"),
        rows=[_row("WO-CZ-SNAPSHOT-NEW", phenomenon="掉漆")],
        run_id="run-cz-snapshot-new",
    )

    assert resolution.topic_id == "CZ-TOPIC-OUTER-HOUSING"
    assert resolution.matched_existing is True
    assert resolution.incremental_supplement_pending is True
    assert resolution.requires_re_review is True


def test_registry_normalizes_cz_snapshot_boundary_values(
    tmp_path: Path,
) -> None:
    audit = AuditStore(tmp_path / "audit.db")
    historical_row = _row(
        "WO-CZ-NORMALIZED-BASE",
        subject="手机屏幕",
        phenomenon="漏液",
    )
    signature = topic_registry_module._topic_signature([historical_row])
    snapshot = CZTopicSnapshot(tmp_path / "cz-topic-snapshot.db")
    snapshot.sync(
        [
            {
                "cz_topic_id": "CZ-TOPIC-NORMALIZED-SNAPSHOT",
                "status": "published",
                "title": "手机屏幕漏液判定",
                "business_line": "自营回收",
                "product_category": "手机",
                "source_topic_key": "TOP-CZ-NORMALIZED-SNAPSHOT",
                "standard_family": _signature_value(
                    signature, "standard_families"
                ),
                "merge_policy": "separate_by_phenomenon",
                "object_key": _signature_value(signature, "object_keys"),
                "phenomenon_value": _signature_value(
                    signature, "phenomenon_values"
                ),
                "query_target": _signature_value(
                    signature, "query_targets"
                ),
                "detection_target": _signature_value(
                    signature, "detection_targets"
                ),
                "platform": _signature_value(signature, "platforms"),
                "brand": _signature_value(signature, "brands"),
                "model_scope": _signature_value(signature, "model_scopes"),
                "threshold_values": _signature_value(
                    signature, "threshold_values"
                ),
            }
        ],
        snapshot_version="local-cz-normalized-v1",
    )

    resolution = TopicRegistry(audit, cz_snapshot=snapshot).integrate(
        proposed_topic_id="TOP-NEW-NORMALIZED-SNAPSHOT",
        topic_key=("direct_mimo", "自营回收", "手机", "cluster-1"),
        rows=[
            _row(
                "WO-CZ-NORMALIZED-NEW",
                subject="手机屏幕",
                phenomenon="漏液",
            )
        ],
        run_id="run-cz-normalized-snapshot-new",
    )

    assert resolution.topic_id == "CZ-TOPIC-NORMALIZED-SNAPSHOT"
    assert resolution.matched_existing is True
    assert resolution.incremental_supplement_pending is True


def test_registry_matches_review_only_local_cz_snapshot_for_manual_review(
    tmp_path: Path,
) -> None:
    audit = AuditStore(tmp_path / "audit.db")
    historical_row = _row("WO-CZ-REVIEW-BASE", phenomenon="碎裂")
    signature = topic_registry_module._topic_signature([historical_row])
    snapshot = CZTopicSnapshot(tmp_path / "cz-topic-snapshot.db")
    snapshot.sync(
        [
            {
                "cz_topic_id": "CZ-TOPIC-REVIEW-ONLY",
                "status": "review",
                "title": "手机外壳外观损伤判定",
                "business_line": "自营回收",
                "product_category": "手机",
                "standard_family": _signature_value(signature, "standard_families"),
                "merge_policy": _signature_value(signature, "merge_policies"),
                "object_key": _signature_value(signature, "object_keys"),
                "phenomenon_value": _signature_value(signature, "phenomenon_values"),
                "query_target": _signature_value(signature, "query_targets"),
                "detection_target": _signature_value(signature, "detection_targets"),
                "platform": _signature_value(signature, "platforms"),
                "brand": _signature_value(signature, "brands"),
                "model_scope": _signature_value(signature, "model_scopes"),
                "threshold_values": _signature_value(signature, "threshold_values"),
            }
        ],
        snapshot_version="local-cz-v1",
    )

    resolution = TopicRegistry(audit, cz_snapshot=snapshot).integrate(
        proposed_topic_id="TOP-NEW-NOT-REVIEW",
        topic_key=("direct_mimo", "自营回收", "手机", "cluster-1"),
        rows=[_row("WO-CZ-REVIEW-NEW", phenomenon="掉漆")],
        run_id="run-cz-review-new",
    )

    assert resolution.topic_id == "TOP-NEW-NOT-REVIEW"
    assert resolution.matched_existing is False
    assert resolution.requires_review is True
    assert resolution.historical_topic_id == "CZ-TOPIC-REVIEW-ONLY"
    assert resolution.decision == "historical_topic_review_required"


def test_workflow_preserves_published_cz_snapshot_topic_id(
    tmp_path: Path,
    monkeypatch,
) -> None:
    row = _row("WO-CZ-WORKFLOW", phenomenon="掉漆")
    signature = topic_registry_module._topic_signature([row])
    snapshot = CZTopicSnapshot(tmp_path / "cz-topic-snapshot.db")
    snapshot.sync(
        [
            {
                "cz_topic_id": "CZ-TOPIC-WORKFLOW",
                "status": "published",
                "title": "手机外壳外观损伤判定",
                "business_line": "自营回收",
                "product_category": "手机",
                "source_topic_key": "TOP-CZ-WORKFLOW",
                "standard_family": _signature_value(signature, "standard_families"),
                "merge_policy": _signature_value(signature, "merge_policies"),
                "object_key": _signature_value(signature, "object_keys"),
                "phenomenon_value": _signature_value(signature, "phenomenon_values"),
                "query_target": _signature_value(signature, "query_targets"),
                "detection_target": _signature_value(signature, "detection_targets"),
                "platform": "通用",
                "brand": "通用",
                "model_scope": "通用",
                "threshold_values": "无明确阈值",
            }
        ],
        snapshot_version="local-cz-workflow-v1",
    )
    monkeypatch.setenv(
        "ANSWER_HUB_CZ_TOPIC_SNAPSHOT_DB",
        str(snapshot.path),
    )

    topics, _mapping, _gaps, _pending = build_topic_review_rows(
        [row],
        use_mimo=False,
        mimo_client=_HighConfidenceDirectMimo(),
        audit_store=AuditStore(tmp_path / "audit.db"),
        run_id="run-cz-workflow",
        clustering_mode="direct_mimo",
        use_standard_references=False,
        enforce_cluster_admission=True,
    )

    assert topics[0]["主题ID"] == "CZ-TOPIC-WORKFLOW"
    assert topics[0]["增量补充状态"] == "待人工复核"
    assert topics[0]["需要重新审核"] == "是"
    assert topics[0]["是否重点复核"] == "是"


def test_registry_stages_new_evidence_for_review_without_mutating_history(
    tmp_path: Path,
) -> None:
    audit = AuditStore(tmp_path / "audit.db")
    registry = TopicRegistry(audit)

    first = registry.integrate(
        proposed_topic_id="TOP-BASE",
        topic_key=("direct_mimo", "自营回收", "手机", "cluster-1"),
        rows=[_row("WO-BASE", phenomenon="碎裂")],
        run_id="run-base",
    )
    supplement = registry.integrate(
        proposed_topic_id="TOP-SUPPLEMENT",
        topic_key=("direct_mimo", "自营回收", "手机", "cluster-2"),
        rows=[_row("WO-SUPPLEMENT", phenomenon="掉漆")],
        run_id="run-supplement",
    )

    assert supplement.topic_id == first.topic_id
    assert supplement.matched_existing is True
    assert supplement.incremental_supplement_pending is True
    assert supplement.requires_re_review is True
    assert supplement.evidence_version == first.evidence_version
    assert supplement.added_member_count == 1
    assert audit.list_registered_topics("自营回收", "手机")[0]["member_count"] == 1
    overlays = audit.list_pending_topic_overlays(first.topic_id)
    assert len(overlays) == 1
    assert overlays[0]["status"] == "pending"

    retried = registry.integrate(
        proposed_topic_id="TOP-SUPPLEMENT-RETRY",
        topic_key=("direct_mimo", "自营回收", "手机", "cluster-2"),
        rows=[_row("WO-SUPPLEMENT", phenomenon="掉漆")],
        run_id="run-supplement-retry",
    )
    assert retried.incremental_supplement_pending is False
    assert retried.added_member_count == 0
    assert len(audit.list_pending_topic_overlays(first.topic_id)) == 1

    approved = audit.review_topic_overlay(overlays[0]["overlay_id"], "approved")
    assert approved["status"] == "approved"
    assert audit.list_registered_topics("自营回收", "手机")[0]["member_count"] == 2


def test_approving_multiple_overlays_preserves_all_boundary_values(
    tmp_path: Path,
) -> None:
    audit = AuditStore(tmp_path / "audit.db")
    registry = TopicRegistry(audit)
    registry.integrate(
        proposed_topic_id="TOP-MULTI-OVERLAY",
        topic_key=("direct_mimo", "自营回收", "手机", "cluster-1"),
        rows=[_row("WO-BASE-MULTI", phenomenon="碎裂")],
        run_id="run-base-multi",
    )
    registry.integrate(
        proposed_topic_id="TOP-MULTI-OVERLAY-1",
        topic_key=("direct_mimo", "自营回收", "手机", "cluster-2"),
        rows=[_row("WO-OVERLAY-1", phenomenon="掉漆")],
        run_id="run-overlay-1",
    )
    registry.integrate(
        proposed_topic_id="TOP-MULTI-OVERLAY-2",
        topic_key=("direct_mimo", "自营回收", "手机", "cluster-3"),
        rows=[_row("WO-OVERLAY-2", phenomenon="凹陷")],
        run_id="run-overlay-2",
    )
    overlays = audit.list_pending_topic_overlays("TOP-MULTI-OVERLAY")
    assert len(overlays) == 2
    for overlay in overlays:
        audit.review_topic_overlay(overlay["overlay_id"], "approved")
    topic = audit.list_registered_topics("自营回收", "手机")[0]
    assert set(topic["signature"]["phenomenon_values"]) == {"碎裂", "掉漆", "凹陷"}


def test_registry_never_merges_different_objects_in_same_standard_family(
    tmp_path: Path,
) -> None:
    registry = TopicRegistry(AuditStore(tmp_path / "audit.db"))

    rear_housing = registry.integrate(
        proposed_topic_id="TOP-REAR-HOUSING",
        topic_key=("direct_mimo", "自营回收", "手机", "cluster-1"),
        rows=[
            _row(
                "WO-REAR-HOUSING",
                subject="手机后壳",
                phenomenon="磕碰",
            )
        ],
        run_id="run-rear-housing",
    )
    frame = registry.integrate(
        proposed_topic_id="TOP-FRAME",
        topic_key=("direct_mimo", "自营回收", "手机", "cluster-2"),
        rows=[
            _row(
                "WO-FRAME",
                subject="手机中框",
                phenomenon="划痕",
            )
        ],
        run_id="run-frame",
    )

    assert frame.topic_id != rear_housing.topic_id
    assert frame.matched_existing is False
    assert {
        row["原始工单ID"]
        for row in frame.rows
    } == {"WO-FRAME"}


def test_historical_similarity_rejects_object_mismatch_before_shared_target() -> None:
    base_signature = {
        "standard_families": [],
        "merge_policies": [],
        "phenomenon_values": [],
        "query_targets": ["phone_housing_appearance"],
        "detection_targets": [],
        "platforms": [],
        "brands": [],
        "model_scopes": [],
        "threshold_values": [],
        "topic_text": "手机外观问题",
    }
    current = {
        **base_signature,
        "object_keys": ["外壳"],
    }
    historical = {
        **base_signature,
        "standard_families": ["手机外壳外观标准"],
        "merge_policies": ["samestandardfamily"],
        "object_keys": ["后置摄像头"],
    }

    similarity, reason = topic_registry_module._signature_similarity(
        current,
        historical,
    )

    assert similarity == 0.0
    assert "对象不同" in reason


def test_historical_similarity_rejects_missing_scope_against_known_scope() -> None:
    base = {
        "standard_families": ["手机外壳外观标准"],
        "merge_policies": ["samestandardfamily"],
        "object_keys": ["外壳"],
        "phenomenon_values": ["碎裂"],
        "query_targets": [],
        "detection_targets": [],
        "brands": [],
        "model_scopes": [],
        "threshold_values": [],
        "topic_text": "手机外壳碎裂如何判定",
    }
    known = {**base, "platforms": ["ios"]}
    missing = {**base, "platforms": []}

    similarity, reason = topic_registry_module._signature_similarity(
        known,
        missing,
    )

    assert similarity == 0.0
    assert "平台" in reason


def test_registry_never_merges_different_product_categories(
    tmp_path: Path,
) -> None:
    registry = TopicRegistry(AuditStore(tmp_path / "audit.db"))

    phone = registry.integrate(
        proposed_topic_id="TOP-PHONE",
        topic_key=("direct_mimo", "自营回收", "手机", "cluster-1"),
        rows=[_row("WO-PHONE", product_type="手机")],
        run_id="run-phone",
    )
    tablet = registry.integrate(
        proposed_topic_id="TOP-TABLET",
        topic_key=("direct_mimo", "自营回收", "平板电脑", "cluster-1"),
        rows=[_row("WO-TABLET", product_type="平板电脑")],
        run_id="run-tablet",
    )

    assert phone.topic_id == "TOP-PHONE"
    assert tablet.topic_id == "TOP-TABLET"
    assert tablet.matched_existing is False


def test_registry_rejects_atomic_evidence_owned_by_another_topic(
    tmp_path: Path,
) -> None:
    audit = AuditStore(tmp_path / "audit.db")
    registry = TopicRegistry(audit)
    row = _row("WO-OWNED-EVIDENCE", product_type="手机")

    registry.integrate(
        proposed_topic_id="TOP-PHONE-OWNER",
        topic_key=("direct_mimo", "自营回收", "手机", "cluster-1"),
        rows=[row],
        run_id="run-phone-owner",
    )
    conflicting_row = dict(row)
    conflicting_row["产品类型"] = "平板电脑"

    with pytest.raises(ValueError, match="已属于其他历史主题"):
        registry.integrate(
            proposed_topic_id="TOP-TABLET-CONFLICT",
            topic_key=("direct_mimo", "自营回收", "平板电脑", "cluster-1"),
            rows=[conflicting_row],
            run_id="run-tablet-conflict",
        )

    assert audit.list_registered_topics("自营回收", "平板电脑") == []


def test_registry_never_merges_different_business_lines(
    tmp_path: Path,
) -> None:
    registry = TopicRegistry(AuditStore(tmp_path / "audit.db"))

    self_operated = registry.integrate(
        proposed_topic_id="TOP-SELF-OPERATED",
        topic_key=("direct_mimo", "自营回收", "手机", "cluster-1"),
        rows=[_row("WO-SELF-OPERATED", business_line="自营回收")],
        run_id="run-self-operated",
    )
    aggregate = registry.integrate(
        proposed_topic_id="TOP-AGGREGATE",
        topic_key=("direct_mimo", "聚合回收", "手机", "cluster-1"),
        rows=[_row("WO-AGGREGATE", business_line="聚合回收")],
        run_id="run-aggregate",
    )

    assert self_operated.topic_id == "TOP-SELF-OPERATED"
    assert aggregate.topic_id == "TOP-AGGREGATE"
    assert aggregate.matched_existing is False


def test_registry_keeps_separate_by_phenomenon_rules_apart(
    tmp_path: Path,
) -> None:
    registry = TopicRegistry(AuditStore(tmp_path / "audit.db"))

    color_spot = registry.integrate(
        proposed_topic_id="TOP-COLOR-SPOT",
        topic_key=("direct_mimo", "自营回收", "手机", "cluster-1"),
        rows=[
            _row(
                "WO-COLOR-SPOT",
                subject="手机屏幕",
                phenomenon="色斑",
            )
        ],
        run_id="run-color-spot",
    )
    dead_pixel = registry.integrate(
        proposed_topic_id="TOP-DEAD-PIXEL",
        topic_key=("direct_mimo", "自营回收", "手机", "cluster-2"),
        rows=[
            _row(
                "WO-DEAD-PIXEL",
                subject="手机屏幕",
                phenomenon="坏点",
            )
        ],
        run_id="run-dead-pixel",
    )

    assert color_spot.topic_id == "TOP-COLOR-SPOT"
    assert dead_pixel.topic_id == "TOP-DEAD-PIXEL"
    assert dead_pixel.matched_existing is False


def test_registry_rejects_mixed_known_and_unknown_product_rows(
    tmp_path: Path,
) -> None:
    registry = TopicRegistry(AuditStore(tmp_path / "audit.db"))

    with pytest.raises(ValueError, match="产品品类"):
        registry.integrate(
            proposed_topic_id="TOP-MIXED-PRODUCT",
            topic_key=("direct_mimo", "自营回收", "手机", "cluster-1"),
            rows=[
                _row("WO-KNOWN", product_type="手机"),
                _row("WO-UNKNOWN", product_type="未知新品类"),
            ],
            run_id="run-mixed-product",
        )


def test_registry_rejects_internal_separate_by_phenomenon_conflict(
    tmp_path: Path,
) -> None:
    registry = TopicRegistry(AuditStore(tmp_path / "audit.db"))

    with pytest.raises(ValueError, match="现象值"):
        registry.integrate(
            proposed_topic_id="TOP-MIXED-SCREEN",
            topic_key=("direct_mimo", "自营回收", "手机", "cluster-1"),
            rows=[
                _row(
                    "WO-COLOR-SPOT-INTERNAL",
                    subject="手机屏幕",
                    phenomenon="色斑",
                ),
                _row(
                    "WO-DEAD-PIXEL-INTERNAL",
                    subject="手机屏幕",
                    phenomenon="坏点",
                ),
            ],
            run_id="run-mixed-screen",
        )


def test_registry_reimport_is_idempotent_for_the_same_atomic_evidence(
    tmp_path: Path,
) -> None:
    registry = TopicRegistry(AuditStore(tmp_path / "audit.db"))
    row = _row("WO-RETRY", phenomenon="磨损")

    first = registry.integrate(
        proposed_topic_id="TOP-RETRY",
        topic_key=("direct_mimo", "自营回收", "手机", "cluster-1"),
        rows=[row],
        run_id="run-first",
    )
    retried = registry.integrate(
        proposed_topic_id="TOP-RETRY-AGAIN",
        topic_key=("direct_mimo", "自营回收", "手机", "cluster-1"),
        rows=[dict(row)],
        run_id="run-retry",
    )

    assert retried.topic_id == first.topic_id
    assert retried.added_member_count == 0
    assert retried.duplicate_member_count == 1
    assert retried.evidence_version == first.evidence_version
    assert len(retried.rows) == 1


def test_registry_reimport_ignores_generic_to_specific_scope_drift(
    tmp_path: Path,
) -> None:
    registry = TopicRegistry(AuditStore(tmp_path / "audit.db"))
    row = _row(
        "WO-SCOPE-DRIFT",
        product_type="单电/微单机身",
        subject="取景器眼罩",
        phenomenon="眼罩缺失",
    )
    row.update(
        {
            "_原子品牌": "通用",
            "_原子机型范围": "通用",
            "_聚类标准族": "相机机身关键部件标准",
            "_聚类合并策略": "separate_by_phenomenon",
            "_聚类现象值": "取景器外观",
        }
    )

    first = registry.integrate(
        proposed_topic_id="TOP-SCOPE-DRIFT-FIRST",
        topic_key=("direct_mimo", "自营回收", "单电/微单机身", "cluster-1"),
        rows=[row],
        run_id="run-scope-drift-first",
    )
    retried_row = dict(row)
    retried_row["_原子品牌"] = "索尼"
    retried_row["_原子机型范围"] = "A6400"
    retried = registry.integrate(
        proposed_topic_id="TOP-SCOPE-DRIFT-SECOND",
        topic_key=("direct_mimo", "自营回收", "单电/微单机身", "cluster-2"),
        rows=[retried_row],
        run_id="run-scope-drift-second",
    )

    assert retried.topic_id == first.topic_id
    assert retried.matched_existing is True
    assert retried.added_member_count == 0
    assert retried.duplicate_member_count == 1


def test_registry_reimport_uses_unchanged_source_before_model_boundary_drift(
    tmp_path: Path,
) -> None:
    registry = TopicRegistry(AuditStore(tmp_path / "audit.db"))
    row = _row(
        "WO-MODEL-BOUNDARY-DRIFT",
        product_type="笔记本",
        subject="屏幕与边框间隙",
        phenomenon="屏幕边缘缝隙或脱胶迹象",
    )
    row["_原子阈值例外"] = "1mm塞规能插入即判定脱胶"

    first = registry.integrate(
        proposed_topic_id="TOP-MODEL-BOUNDARY-FIRST",
        topic_key=("direct_mimo", "自营回收", "笔记本", "cluster-1"),
        rows=[row],
        run_id="run-model-boundary-first",
    )
    retried_row = dict(row)
    retried_row["对象/部位"] = "屏幕边缘"
    retried_row["异常现象"] = "缝隙/脱胶迹象"
    retried_row["核心问题"] = "笔记本屏幕边缘缝隙和脱胶迹象如何判定"
    retried_row["_原子阈值例外"] = "塞规能插入则判定为屏幕脱胶"
    retried = registry.integrate(
        proposed_topic_id="TOP-MODEL-BOUNDARY-SECOND",
        topic_key=("direct_mimo", "自营回收", "笔记本", "cluster-2"),
        rows=[retried_row],
        run_id="run-model-boundary-second",
    )

    assert retried.topic_id == first.topic_id
    assert retried.matched_existing is True
    assert retried.added_member_count == 0
    assert retried.duplicate_member_count == 1


def test_registry_reimport_rejects_changed_atomic_boundary(
    tmp_path: Path,
) -> None:
    registry = TopicRegistry(AuditStore(tmp_path / "audit.db"))
    row = _row("WO-CHANGED-BOUNDARY", phenomenon="碎裂")
    row["_原子品牌"] = "Apple"
    row["_原子机型范围"] = "iPhone 15"

    registry.integrate(
        proposed_topic_id="TOP-CHANGED-BOUNDARY-FIRST",
        topic_key=("direct_mimo", "自营回收", "手机", "cluster-1"),
        rows=[row],
        run_id="run-changed-boundary-first",
    )
    changed = dict(row)
    changed["_原子机型范围"] = "iPhone 16"
    changed["聊天内容"] = "同一工单后来确认实际设备为iPhone 16。"

    with pytest.raises(ValueError, match="已属于其他历史主题"):
        registry.integrate(
            proposed_topic_id="TOP-CHANGED-BOUNDARY-SECOND",
            topic_key=("direct_mimo", "自营回收", "手机", "cluster-2"),
            rows=[changed],
            run_id="run-changed-boundary-second",
        )


def test_registry_does_not_remerge_different_atomics_from_same_work_order(
    tmp_path: Path,
) -> None:
    registry = TopicRegistry(AuditStore(tmp_path / "audit.db"))
    first_row = _row("WO-MULTI", phenomenon="碎裂")
    first_row["_原子知识ID"] = "WO-MULTI-U1"
    second_row = _row("WO-MULTI", phenomenon="掉漆")
    second_row["_原子知识ID"] = "WO-MULTI-U2"

    first = registry.integrate(
        proposed_topic_id="TOP-MULTI-1",
        topic_key=("direct_mimo", "自营回收", "手机", "cluster-1"),
        rows=[first_row],
        run_id="run-multi-1",
    )
    second = registry.integrate(
        proposed_topic_id="TOP-MULTI-2",
        topic_key=("direct_mimo", "自营回收", "手机", "cluster-2"),
        rows=[second_row],
        run_id="run-multi-2",
    )

    assert second.topic_id != first.topic_id
    assert second.matched_existing is False


def test_registry_respects_explicit_model_scope_differences(
    tmp_path: Path,
) -> None:
    registry = TopicRegistry(AuditStore(tmp_path / "audit.db"))
    first_row = _row("WO-MODEL-15", phenomenon="碎裂")
    first_row["_原子机型范围"] = "iPhone 15"
    second_row = _row("WO-MODEL-16", phenomenon="掉漆")
    second_row["_原子机型范围"] = "iPhone 16"

    first = registry.integrate(
        proposed_topic_id="TOP-MODEL-15",
        topic_key=("direct_mimo", "自营回收", "手机", "cluster-1"),
        rows=[first_row],
        run_id="run-model-15",
    )
    second = registry.integrate(
        proposed_topic_id="TOP-MODEL-16",
        topic_key=("direct_mimo", "自营回收", "手机", "cluster-2"),
        rows=[second_row],
        run_id="run-model-16",
    )

    assert second.topic_id != first.topic_id
    assert second.matched_existing is False


def test_registry_rejects_modified_original_work_order_id(
    tmp_path: Path,
) -> None:
    registry = TopicRegistry(AuditStore(tmp_path / "audit.db"))
    row = _row("WO-ORIGINAL")
    row["工单ID"] = "WO-MODIFIED"

    with pytest.raises(ValueError, match="原始工单ID"):
        registry.integrate(
            proposed_topic_id="TOP-WORK-ORDER-MISMATCH",
            topic_key=("direct_mimo", "自营回收", "手机", "cluster-1"),
            rows=[row],
            run_id="run-work-order-mismatch",
        )


def test_workflow_reuses_topic_id_and_accumulates_original_work_orders(
    tmp_path: Path,
) -> None:
    audit = AuditStore(tmp_path / "audit.db")

    first_topics, _first_mapping, _first_gaps, _first_pending = (
        build_topic_review_rows(
            [_row("WO-WORKFLOW-001", phenomenon="碎裂")],
            use_mimo=False,
            mimo_client=_HighConfidenceDirectMimo(),
            audit_store=audit,
            run_id="run-workflow-1",
            clustering_mode="direct_mimo",
            use_standard_references=False,
            enforce_cluster_admission=True,
        )
    )
    second_topics, second_mapping, _second_gaps, _second_pending = (
        build_topic_review_rows(
            [_row("WO-WORKFLOW-002", phenomenon="掉漆")],
            use_mimo=False,
            mimo_client=_HighConfidenceDirectMimo(),
            audit_store=audit,
            run_id="run-workflow-2",
            clustering_mode="direct_mimo",
            use_standard_references=False,
            enforce_cluster_admission=True,
        )
    )

    assert second_topics[0]["主题ID"] == first_topics[0]["主题ID"]
    assert second_topics[0]["增量补充状态"] == "待人工复核"
    assert second_topics[0]["需要重新审核"] == "是"
    assert second_topics[0]["是否重点复核"] == "是"
    assert second_topics[0]["主题样本数"] == 2
    assert set(second_topics[0]["主题工单ID"].splitlines()) == {
        "WO-WORKFLOW-001",
        "WO-WORKFLOW-002",
    }
    assert {
        row["原始工单ID"]
        for row in second_mapping
    } == {"WO-WORKFLOW-001", "WO-WORKFLOW-002"}


def test_cluster_only_reuses_topic_id_across_runs_without_transcription(
    tmp_path: Path,
) -> None:
    audit = AuditStore(tmp_path / "audit.db")

    first_topics, _first_mapping, _first_gaps, _first_pending = (
        build_topic_review_rows(
            [_row("WO-CLUSTER-ONLY-001", phenomenon="碎裂")],
            use_mimo=False,
            mimo_client=_HighConfidenceDirectMimo(),
            audit_store=audit,
            run_id="run-cluster-only-1",
            clustering_mode="direct_mimo",
            use_standard_references=False,
            cluster_only=True,
            enforce_cluster_admission=True,
        )
    )
    second_topics, _second_mapping, _second_gaps, _second_pending = (
        build_topic_review_rows(
            [_row("WO-CLUSTER-ONLY-002", phenomenon="掉漆")],
            use_mimo=False,
            mimo_client=_HighConfidenceDirectMimo(),
            audit_store=audit,
            run_id="run-cluster-only-2",
            clustering_mode="direct_mimo",
            use_standard_references=False,
            cluster_only=True,
            enforce_cluster_admission=True,
        )
    )

    assert second_topics[0]["聚类主题ID"] == first_topics[0]["聚类主题ID"]
    assert second_topics[0]["历史主题处理结果"] == "incremental_supplement_pending"
    assert second_topics[0]["增量补充状态"] == "待人工复核"
    assert second_topics[0]["需要重新审核"] == "是"
    assert second_topics[0]["是否重点复核"] == "是"
    assert second_topics[0]["主题样本数"] == 2
    assert set(second_topics[0]["主题工单ID"].splitlines()) == {
        "WO-CLUSTER-ONLY-001",
        "WO-CLUSTER-ONLY-002",
    }


def test_registry_reuses_first_round_topic_for_cross_work_order_near_synonym(
    tmp_path: Path,
) -> None:
    audit = AuditStore(tmp_path / "audit.db")

    first_row = _row(
        "WO-CROSS-NEAR-001",
        subject="手机外壳",
        phenomenon="碎裂",
    )
    second_row = _row(
        "WO-CROSS-NEAR-002",
        subject="手机外壳",
        phenomenon="裂纹",
    )

    registry = TopicRegistry(audit)
    first = registry.integrate(
        proposed_topic_id="TOP-CROSS-NEAR-001",
        topic_key=("direct_mimo", "自营回收", "手机", "round1-cluster-1"),
        rows=[first_row],
        run_id="run-cross-near-1",
    )
    second = registry.integrate(
        proposed_topic_id="TOP-CROSS-NEAR-002",
        topic_key=("direct_mimo", "自营回收", "手机", "round2-cluster-1"),
        rows=[second_row],
        run_id="run-cross-near-2",
    )

    assert first.topic_id == "TOP-CROSS-NEAR-001"
    assert first.matched_existing is False
    assert second.topic_id == first.topic_id
    assert second.historical_topic_id == first.topic_id
    assert second.matched_existing is True
    assert second.decision == "incremental_supplement_pending"
    assert second.incremental_supplement_pending is True
    assert second.requires_re_review is True
    assert second.added_member_count == 1
    assert second.duplicate_member_count == 0
    assert second.reason == "新增非重复证据已隔离，等待人工复核后更新可信历史主题"
    assert len(audit.list_registered_topics("自营回收", "手机")) == 1
    overlays = audit.list_pending_topic_overlays(first.topic_id)
    assert len(overlays) == 1
    assert overlays[0]["added_member_count"] == 1
    assert overlays[0]["duplicate_member_count"] == 0


def test_cluster_only_writer_passes_incremental_registry_and_exports_audit_fields(
    tmp_path: Path,
) -> None:
    audit = AuditStore(tmp_path / "audit.db")

    first_summary = write_cluster_only_workbook(
        [_row("WO-CLUSTER-WRITER-001", phenomenon="碎裂")],
        tmp_path / "first.xlsx",
        use_mimo=False,
        mimo_client=_HighConfidenceDirectMimo(),
        clustering_mode="direct_mimo",
        audit_store=audit,
        run_id="run-cluster-writer-1",
        enforce_cluster_admission=True,
    )
    second_summary = write_cluster_only_workbook(
        [_row("WO-CLUSTER-WRITER-002", phenomenon="掉漆")],
        tmp_path / "second.xlsx",
        use_mimo=False,
        mimo_client=_HighConfidenceDirectMimo(),
        clustering_mode="direct_mimo",
        audit_store=audit,
        run_id="run-cluster-writer-2",
        enforce_cluster_admission=True,
    )

    assert first_summary["incremental_created_topics"] == 1
    assert second_summary["incremental_merged_topics"] == 1
    assert second_summary["incremental_re_review_topics"] == 1
    workbook = load_workbook(tmp_path / "second.xlsx", read_only=True, data_only=True)
    worksheet = workbook["聚类结果"]
    rows = list(worksheet.iter_rows(values_only=True))
    workbook.close()
    exported = dict(zip(rows[0], rows[1]))
    assert exported["历史主题处理结果"] == "incremental_supplement_pending"
    assert exported["增量补充状态"] == "待人工复核"
    assert exported["需要重新审核"] == "是"


def test_cluster_only_does_not_register_review_required_topic(
    tmp_path: Path,
) -> None:
    audit = AuditStore(tmp_path / "audit.db")

    topics, _mapping, _gaps, _pending = build_topic_review_rows(
        [_row("WO-CLUSTER-REVIEW-001")],
        use_mimo=False,
        mimo_client=_ReviewRequiredDirectMimo(),
        audit_store=audit,
        run_id="run-cluster-review-1",
        clustering_mode="direct_mimo",
        use_standard_references=False,
        cluster_only=True,
        enforce_cluster_admission=True,
    )

    assert topics[0]["聚类准入状态"] == "待人工聚类复核"
    assert topics[0]["是否重点复核"] == "是"
    assert audit.list_registered_topics("自营回收", "手机") == []


def test_workflow_emits_one_candidate_when_two_groups_resolve_same_topic(
    tmp_path: Path,
    monkeypatch,
) -> None:
    audit = AuditStore(tmp_path / "audit.db")
    first = _row("WO-SAME-RUN-001", phenomenon="碎裂")
    second = _row("WO-SAME-RUN-002", phenomenon="掉漆")
    for row in (first, second):
        row.update(
            {
                "语义标注状态": "topic_signal_labeled",
                "语义标注置信度": 0.95,
                "_聚类裁决提供方": "mimo-direct",
                "_聚类裁决置信度": 0.95,
                "_聚类决策": "模型标签一致合并",
                "_聚类需要复核": False,
                "_原子需要复核": False,
            }
        )

    def split_groups(*_args, **_kwargs):
        return [
            (("direct_mimo", "self", "group-1"), [first]),
            (("direct_mimo", "self", "group-2"), [second]),
        ], {
            "requested_mode": "direct_mimo",
            "effective_mode": "direct_mimo",
            "provider": "mimo-atomic-extraction+direct-topic-clustering",
            "cluster_count": 2,
        }

    monkeypatch.setattr(
        workflow_module,
        "_direct_mimo_topic_groups",
        split_groups,
    )

    topics, mapping, gaps, pending = build_topic_review_rows(
        [first, second],
        use_mimo=False,
        mimo_client=SimpleNamespace(
            config=SimpleNamespace(model="same-run-merge-test")
        ),
        audit_store=audit,
        run_id="run-same-run-merge",
        clustering_mode="direct_mimo",
        use_standard_references=False,
        enforce_cluster_admission=True,
    )

    assert not gaps
    assert not pending
    assert len(topics) == 1
    assert len({topic["主题ID"] for topic in topics}) == 1
    assert topics[0]["主题样本数"] == 2
    assert {
        row["原始工单ID"]
        for row in mapping
    } == {"WO-SAME-RUN-001", "WO-SAME-RUN-002"}


def test_rule_fallback_coalesces_same_boundary_with_different_labels() -> None:
    first = _row("WO-RULE-BOUNDARY-001", phenomenon="碎裂")
    second = _row("WO-RULE-BOUNDARY-002", phenomenon="掉漆")
    second.update(
        {
            "模型主题二级分类": "外观损伤",
            "主标准路径": "外观损伤处理路径",
            "核心问题": "手机外壳掉色应该怎样判断",
            "聊天内容": "外壳漆面脱落时按哪种口径处理？",
        }
    )

    topics, mapping, gaps, pending = build_topic_review_rows(
        [first, second],
        use_mimo=False,
        clustering_mode="rule",
        use_standard_references=False,
    )

    assert len(topics) == 1
    assert topics[0]["主题样本数"] == 2
    assert len(mapping) == 2
    assert not gaps
    assert not pending


def test_workflow_does_not_write_history_when_admission_is_disabled(
    tmp_path: Path,
) -> None:
    audit = AuditStore(tmp_path / "audit.db")

    build_topic_review_rows(
        [_row("WO-NO-ADMISSION")],
        use_mimo=False,
        audit_store=audit,
        run_id="run-no-admission",
        clustering_mode="rule",
        use_standard_references=False,
        enforce_cluster_admission=False,
    )

    assert audit.list_registered_topics("自营回收", "手机") == []


def test_workflow_cluster_admission_rejects_mixed_unknown_product(
    tmp_path: Path,
) -> None:
    audit = AuditStore(tmp_path / "audit.db")

    topics, _mapping, _gaps, pending = build_topic_review_rows(
        [
            _row("WO-ADMISSION-KNOWN", product_type="手机"),
            _row("WO-ADMISSION-UNKNOWN", product_type="未知新品类"),
        ],
        use_mimo=False,
        audit_store=audit,
        run_id="run-admission-mixed-product",
        clustering_mode="rule",
        use_standard_references=False,
        enforce_cluster_admission=True,
    )

    assert len(topics) == 1
    assert topics[0]["主题状态"] == "provisional_singleton_review_pending"
    assert topics[0]["主题工单ID"] == "WO-ADMISSION-KNOWN"
    assert len(pending) == 1
    unknown_pending = next(
        row
        for row in pending
        if row["原始工单ID"] == "WO-ADMISSION-UNKNOWN"
    )
    assert "产品品类" in unknown_pending["待聚合原因"]


def test_registry_bootstraps_existing_audit_candidates_before_matching(
    tmp_path: Path,
) -> None:
    audit = AuditStore(tmp_path / "audit.db")
    audit.save_candidate(
        model_run_id="legacy-model-run",
        run_id="legacy-run",
        record_id="TOP-LEGACY",
        candidate={
            "主题ID": "TOP-LEGACY",
            "回收业务层级": "自营回收",
            "适用范围": "手机",
            "聚类准入状态": "已自动放行",
            "聚类准入置信度": 0.92,
            "主题问题意图": "标准判定",
            "主题对象/部位": "手机外壳",
            "主题异常现象": "碎裂",
            "主题解题方式": "按外观质检口径判定",
            "主题事实证据包": json.dumps(
                {
                    "facts": [
                        {
                            "source_record_id": "LEGACY-RECORD",
                            "work_order_id": "WO-LEGACY",
                            "atomic_question": "手机外壳碎裂如何判定",
                            "human_core_problem": "手机外壳碎裂如何判定",
                            "human_judgment_conclusion": "按外壳外观口径判定",
                            "conversation_excerpt": "手机外壳碎裂应该怎么判定？",
                            "historical_actual_reply": "",
                            "judgment_basis": "",
                            "image_urls": [],
                            "image_usable": False,
                        }
                    ]
                },
                ensure_ascii=False,
            ),
        },
    )

    resolved = TopicRegistry(audit).integrate(
        proposed_topic_id="TOP-NEW",
        topic_key=("direct_mimo", "自营回收", "手机", "cluster-1"),
        rows=[_row("WO-NEW", phenomenon="掉漆")],
        run_id="run-new",
    )

    assert resolved.topic_id == "TOP-LEGACY"
    assert resolved.matched_existing is True
    assert {
        row["原始工单ID"]
        for row in resolved.rows
    } == {"WO-LEGACY", "WO-NEW"}


def test_registry_does_not_bootstrap_unverified_review_pending_candidate(
    tmp_path: Path,
) -> None:
    audit = AuditStore(tmp_path / "audit.db")
    audit.save_candidate(
        model_run_id="legacy-unverified-model-run",
        run_id="legacy-unverified-run",
        record_id="TOP-LEGACY-UNVERIFIED",
        candidate={
            "主题ID": "TOP-LEGACY-UNVERIFIED",
            "回收业务层级": "自营回收",
            "适用范围": "手机",
            "主题问题意图": "标准判定",
            "主题对象/部位": "手机外壳",
            "主题异常现象": "碎裂",
            "主题解题方式": "按外观质检口径判定",
            "主题事实证据包": json.dumps(
                {
                    "facts": [
                        {
                            "source_record_id": "LEGACY-UNVERIFIED",
                            "work_order_id": "WO-LEGACY-UNVERIFIED",
                            "atomic_question": "手机外壳碎裂如何判定",
                            "human_core_problem": "手机外壳碎裂如何判定",
                            "human_judgment_conclusion": "待人工确认",
                            "conversation_excerpt": "手机外壳碎裂应该怎么判定？",
                            "image_urls": [],
                        }
                    ]
                },
                ensure_ascii=False,
            ),
        },
    )

    resolved = TopicRegistry(audit).integrate(
        proposed_topic_id="TOP-VERIFIED-NEW",
        topic_key=("direct_mimo", "自营回收", "手机", "cluster-1"),
        rows=[_row("WO-VERIFIED-NEW", phenomenon="掉漆")],
        run_id="run-verified-new",
    )

    assert resolved.topic_id == "TOP-VERIFIED-NEW"
    assert resolved.matched_existing is False


def test_registry_updates_signature_after_new_compatible_evidence(
    tmp_path: Path,
) -> None:
    audit = AuditStore(tmp_path / "audit.db")
    registry = TopicRegistry(audit)

    registry.integrate(
        proposed_topic_id="TOP-SIGNATURE",
        topic_key=("direct_mimo", "自营回收", "手机", "cluster-1"),
        rows=[_row("WO-SIGNATURE-1", phenomenon="碎裂")],
        run_id="run-signature-1",
    )
    registry.integrate(
        proposed_topic_id="TOP-SIGNATURE-NEW",
        topic_key=("direct_mimo", "自营回收", "手机", "cluster-2"),
        rows=[_row("WO-SIGNATURE-2", phenomenon="掉漆")],
        run_id="run-signature-2",
    )

    overlay = audit.list_pending_topic_overlays("TOP-SIGNATURE")[0]
    audit.review_topic_overlay(overlay["overlay_id"], "approved")
    stored = audit.list_registered_topics("自营回收", "手机")[0]
    assert set(stored["signature"]["phenomenon_values"]) == {
        "碎裂",
        "掉漆",
    }


def test_registry_serializes_concurrent_compatible_topic_creation(
    tmp_path: Path,
) -> None:
    audit = AuditStore(tmp_path / "audit.db")
    registries = [TopicRegistry(audit) for _index in range(8)]
    barrier = threading.Barrier(len(registries))

    def integrate(index: int) -> str:
        barrier.wait()
        resolution = registries[index].integrate(
            proposed_topic_id=f"TOP-CONCURRENT-{index}",
            topic_key=(
                "direct_mimo",
                "自营回收",
                "手机",
                f"cluster-{index}",
            ),
            rows=[
                _row(
                    f"WO-CONCURRENT-{index}",
                    phenomenon="碎裂" if index % 2 == 0 else "掉漆",
                )
            ],
            run_id=f"run-concurrent-{index}",
        )
        return resolution.topic_id

    with ThreadPoolExecutor(max_workers=len(registries)) as executor:
        topic_ids = list(executor.map(integrate, range(len(registries))))

    assert len(set(topic_ids)) == 1
    topics = audit.list_registered_topics("自营回收", "手机")
    assert len(topics) == 1
    overlays = audit.list_pending_topic_overlays(topics[0]["topic_id"])
    for overlay in overlays:
        audit.review_topic_overlay(overlay["overlay_id"], "approved")
    topics = audit.list_registered_topics("自营回收", "手机")
    assert topics[0]["member_count"] == len(registries)


def test_registry_does_not_merge_unrelated_topics_on_proposed_id_collision(
    tmp_path: Path,
) -> None:
    registry = TopicRegistry(AuditStore(tmp_path / "audit.db"))
    speaker = _row(
        "WO-SPEAKER",
        subject="扬声器",
        phenomenon="声音小",
    )
    bluetooth = _row(
        "WO-BLUETOOTH",
        subject="蓝牙",
        phenomenon="连接失败",
    )

    first = registry.integrate(
        proposed_topic_id="TOP-COLLISION",
        topic_key=("rule", "自营回收", "手机", "shared-key"),
        rows=[speaker],
        run_id="run-speaker",
    )
    second = registry.integrate(
        proposed_topic_id="TOP-COLLISION",
        topic_key=("rule", "自营回收", "手机", "shared-key"),
        rows=[bluetooth],
        run_id="run-bluetooth",
    )

    assert first.topic_id == "TOP-COLLISION"
    assert second.topic_id != first.topic_id
    assert second.matched_existing is False
    assert {
        row["原始工单ID"]
        for row in second.rows
    } == {"WO-BLUETOOTH"}


def test_registry_routes_uncertain_historical_match_to_manual_review(
    tmp_path: Path,
) -> None:
    audit = AuditStore(tmp_path / "audit.db")
    registry = TopicRegistry(audit)
    first_row = _row(
        "WO-SPEAKER-OLD",
        subject="扬声器",
        phenomenon="声音偏小",
    )
    first_row["核心问题"] = "扬声器声音偏小如何判断"
    first_row["聊天内容"] = "扬声器声音偏小如何判断？"
    first_row["模型主题一级分类"] = "功能问题"
    first_row["模型主题二级分类"] = "功能异常"
    first_row["问题意图"] = "检测核验"
    first_row["解题方式"] = "现场功能测试"
    second_row = _row(
        "WO-SPEAKER-NEW",
        subject="扬声器",
        phenomenon="声音很小",
    )
    second_row["核心问题"] = "扬声器声音很小怎么检测"
    second_row["聊天内容"] = "扬声器声音很小怎么检测？"
    second_row["模型主题一级分类"] = "功能问题"
    second_row["模型主题二级分类"] = "功能异常"
    second_row["问题意图"] = "检测核验"
    second_row["解题方式"] = "现场功能测试"

    registry.integrate(
        proposed_topic_id="TOP-SPEAKER-OLD",
        topic_key=("direct_mimo", "自营回收", "手机", "cluster-1"),
        rows=[first_row],
        run_id="run-speaker-old",
    )
    uncertain = registry.integrate(
        proposed_topic_id="TOP-SPEAKER-NEW",
        topic_key=("direct_mimo", "自营回收", "手机", "cluster-2"),
        rows=[second_row],
        run_id="run-speaker-new",
    )

    assert uncertain.requires_review is True
    assert uncertain.decision == "historical_topic_review_required"
    assert uncertain.added_member_count == 0
    assert len(
        audit.list_registered_topics("自营回收", "手机")
    ) == 1


def test_workflow_stops_uncertain_historical_match_before_transcription(
    tmp_path: Path,
) -> None:
    audit = AuditStore(tmp_path / "audit.db")
    old_row = _row(
        "WO-WORKFLOW-SPEAKER-OLD",
        subject="扬声器",
        phenomenon="声音偏小",
    )
    old_row.update(
        {
            "核心问题": "扬声器声音偏小如何判断",
            "聊天内容": "扬声器声音偏小如何判断？",
            "模型主题一级分类": "功能问题",
            "模型主题二级分类": "功能异常",
            "问题意图": "检测核验",
            "解题方式": "现场功能测试",
        }
    )
    new_row = _row(
        "WO-WORKFLOW-SPEAKER-NEW",
        subject="扬声器",
        phenomenon="声音很小",
    )
    new_row.update(
        {
            "核心问题": "扬声器声音很小怎么检测",
            "聊天内容": "扬声器声音很小怎么检测？",
            "模型主题一级分类": "功能问题",
            "模型主题二级分类": "功能异常",
            "问题意图": "检测核验",
            "解题方式": "现场功能测试",
        }
    )

    build_topic_review_rows(
        [old_row],
        use_mimo=False,
        mimo_client=_HighConfidenceDirectMimo(),
        audit_store=audit,
        run_id="run-workflow-speaker-old",
        clustering_mode="direct_mimo",
        use_standard_references=False,
        enforce_cluster_admission=True,
    )
    topics, _mapping, _gaps, pending = build_topic_review_rows(
        [new_row],
        use_mimo=False,
        mimo_client=_HighConfidenceDirectMimo(),
        audit_store=audit,
        run_id="run-workflow-speaker-new",
        clustering_mode="direct_mimo",
        use_standard_references=False,
        enforce_cluster_admission=True,
    )

    assert topics == []
    assert len(pending) == 1
    assert pending[0]["待聚合状态"] == "pending_historical_topic_review"
    assert "边界不完整" in pending[0]["待聚合原因"]
